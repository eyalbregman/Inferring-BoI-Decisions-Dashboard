# -*- coding: utf-8 -*-
"""results/hawkish_dovish_fed_anchored.xlsx

Sheet 1  Lexicon (Fed-anchored) : the hawk/dove terms, their category, direction and source
Sheet 2  Announcement scores    : one row per raw BoI announcement, score = (H-D)/(H+D)
Sheet 3  Method
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'results', 'hawkish_dovish_fed_anchored.xlsx')

import sys
sys.path.insert(0, os.path.join(ROOT, 'scripts'))
from lexicon_fed import LEXICON

scores = json.load(open(os.path.join(ROOT, 'results', 'fed_lexicon_scores.json'), encoding='utf-8'))
hits = json.load(open(os.path.join(ROOT, 'results', 'fed_lexicon_hits.json'), encoding='utf-8'))
cover = hits['coverage']

BASE = Font(name="Arial", size=10)
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F4E79")
FILL_H = PatternFill("solid", fgColor="FBE4E1")
FILL_D = PatternFill("solid", fgColor="DEEAF2")
FILL_C = PatternFill("solid", fgColor="F2F2F2")
POS = PatternFill("solid", fgColor="F4C7C3")
NEG = PatternFill("solid", fgColor="BDD7EE")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def head(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADF; cell.fill = HEADFILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    ws.row_dimensions[1].height = 30

DIRNAME = {"H": "Hawkish", "D": "Dovish", "C": "context"}
SRCNAME = {"fed": "Fed contrast", "fed+judg": "Fed contrast + judgment", "judgment": "judgment"}
DORDER = {"H": 0, "D": 1, "C": 2}

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Sheet 1: Lexicon
ws = wb.active
ws.title = "Lexicon (Fed-anchored)"
ws.append(["Term", "Category", "Direction", "Source", "Why it reads that way",
           "Example from the 62 BoI announcements", "Occurrences", "In N of 62"])
lex = sorted(LEXICON, key=lambda e: (e[1].lower(), DORDER[e[2]], e[0].lower()))
for term, cat, d, src, logic in lex:
    tot, docf, ex = cover.get(term, (0, 0, ""))
    ws.append([term, cat, DIRNAME[d], SRCNAME[src], logic, ex, tot, docf])
for i, (term, cat, d, src, logic) in enumerate(lex, start=2):
    for c in range(1, 9):
        cell = ws.cell(row=i, column=c); cell.font = BASE; cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (1, 2, 5, 6)))
    dc = ws.cell(row=i, column=3)
    dc.fill = {"H": FILL_H, "D": FILL_D, "C": FILL_C}[d]
    dc.font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=i, column=7).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=i, column=8).alignment = Alignment(horizontal="center", vertical="top")
for i, w in enumerate([32, 24, 11, 22, 50, 60, 12, 11], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
head(ws, 8)

# ---------------------------------------------------------------- Sheet 2: scores
ws2 = wb.create_sheet("Announcement scores")
ws2.append(["Date", "Decision (reference only)", "Hawkish hits (H)", "Dovish hits (D)",
            "Score  (H - D) / (H + D)", "Calculation",
            "Hawkish per 1,000 words", "Dovish per 1,000 words"])
for i, r in enumerate(scores, start=2):
    ws2.append([r['date'], r['decision'], r['H'], r['D'], round(r['score'], 4),
                f"({r['H']} - {r['D']}) / ({r['H']} + {r['D']}) = {r['score']:+.3f}",
                r['hawk_per_1k'], r['dove_per_1k']])
    for c in range(1, 9):
        cell = ws2.cell(row=i, column=c); cell.font = BASE; cell.border = BORDER
        if c in (3, 4, 7, 8):
            cell.alignment = Alignment(horizontal="center")
    sc = ws2.cell(row=i, column=5)
    sc.number_format = "+0.000;-0.000;0.000"
    sc.font = Font(name="Arial", size=10, bold=True)
    sc.fill = POS if r['score'] > 0 else (NEG if r['score'] < 0 else FILL_C)
    ws2.cell(row=i, column=3).fill = FILL_H
    ws2.cell(row=i, column=4).fill = FILL_D
for i, w in enumerate([12, 22, 15, 14, 16, 30, 16, 16], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
head(ws2, 8)
b = len(scores) + 3
sH = sum(r['H'] for r in scores); sD = sum(r['D'] for r in scores)
ws2.cell(row=b, column=1, value="All announcements").font = Font(name="Arial", size=10, bold=True)
ws2.cell(row=b, column=3, value=sH).font = Font(name="Arial", size=10, bold=True)
ws2.cell(row=b, column=4, value=sD).font = Font(name="Arial", size=10, bold=True)
sc = ws2.cell(row=b, column=5, value=round((sH - sD) / (sH + sD), 4))
sc.font = Font(name="Arial", size=10, bold=True); sc.number_format = "+0.000;-0.000;0.000"
ws2.cell(row=b, column=6, value=f"({sH} - {sD}) / ({sH} + {sD}) = {(sH-sD)/(sH+sD):+.3f}").font = BASE
for j, (lab, dec) in enumerate([("mean score, hikes", "raise"), ("mean score, holds", "maintain"),
                                ("mean score, cuts", "lower")], start=1):
    v = [r['score'] for r in scores if r['decision'] == dec]
    ws2.cell(row=b + j, column=1, value=lab).font = BASE
    mc = ws2.cell(row=b + j, column=5, value=round(sum(v) / len(v), 4))
    mc.font = BASE; mc.number_format = "+0.000;-0.000;0.000"

# ---------------------------------------------------------------- Sheet 3: Method
ws3 = wb.create_sheet("Method")
nH = sum(1 for e in LEXICON if e[2] == 'H'); nD = sum(1 for e in LEXICON if e[2] == 'D')
meth = [
    ("Fed-anchored hawkish / dovish lexicon", True),
    ("", False),
    ("The idea (after Doh, Song & Yang 2025, 'Deciphering Federal Reserve Communication')", True),
    ("For every FOMC meeting the Federal Reserve Board staff drafts a DOVISH alternative", False),
    ("statement (Alternative A) and a HAWKISH one (Alternative C / D). The released statement", False),
    ("sits somewhere between them. Doh-Song-Yang score a statement's tone by how close its", False),
    ("wording is to the hawkish pole vs. the dovish pole. The Bank of Israel publishes no such", False),
    ("drafts, so we borrow the Fed's: we use the Fed's own extreme drafts to learn which words", False),
    ("and phrasings are hawkish vs. dovish, then apply that vocabulary to the BoI announcements.", False),
    ("", False),
    ("How the lexicon was built", True),
    ("1. Downloaded 136 Fed Bluebook / Tealbook-B PDFs (federalreserve.gov, 2004-2020) and", False),
    ("   extracted the Alt A / B / C / D draft statements for 75 meetings (2009-2020).", False),
    ("2. Pooled all Alt A text (dovish pole) vs. all Alt C/D text (hawkish pole) and computed", False),
    ("   weighted log-odds (Monroe, Colaresi & Quinn 2008) to find the discriminating 1-3 grams.", False),
    ("3. Hand-read a spanning sample of Alt A vs. Alt C pairs (liftoff 2015-18, cuts 2019-20).", False),
    ("4. Curated the result with judgment, kept only terms that also occur in the BoI corpus,", False),
    ("   translated Fed phrasings into BoI idiom, and added the decision + forward-guidance", False),
    ("   vocabulary ('raise / restrictive / tight' hawkish; 'lower / accommodative / patient'", False),
    ("   dovish) so that, per the brief, the decision and guidance themselves move the score.", False),
    ("", False),
    (f"Result: {len(LEXICON)} terms  ({nH} hawkish, {nD} dovish).  'Source' column: 'Fed contrast'", False),
    ("= straight from the Alt-A/Alt-C comparison; 'Fed contrast + judgment' = generalised;", False),
    ("'judgment' = added by hand (BoI idiom, decision vocabulary).", False),
    ("", False),
    ("Scoring the announcements", True),
    ("Each RAW BoI announcement (headline, decision sentence and all stance language KEPT;", False),
    ("navigation, dates, figure captions and Hebrew removed; the 4x-repeated headline collapsed", False),
    ("to one). Every lexicon term is matched; overlapping matches are counted once (longer wins).", False),
    ("   H = hawkish hits,  D = dovish hits", False),
    ("   score = (H - D) / (H + D)      in [-1, +1]      shown per row in the 'Calculation' column", False),
    ("+1 = every stance/description cue is hawkish; -1 = every cue is dovish; 0 = balanced.", False),
    ("", False),
    ("What it shows", True),
    ("Hikes score clearly hawkish (mean %+.2f); holds (%+.2f) and cuts (%+.2f) both score dovish."
     % (sum(r['score'] for r in scores if r['decision'] == 'raise') / max(1, sum(1 for r in scores if r['decision'] == 'raise')),
        sum(r['score'] for r in scores if r['decision'] == 'maintain') / max(1, sum(1 for r in scores if r['decision'] == 'maintain')),
        sum(r['score'] for r in scores if r['decision'] == 'lower') / max(1, sum(1 for r in scores if r['decision'] == 'lower'))), False),
    ("As with every other method in this project, the lexicon separates hikes from the rest but", False),
    ("barely separates cuts from holds: the BoI writes cuts in the same measured register as holds.", False),
    ("", False),
    ("Caveats", True),
    ("- Transfer method: the vocabulary is the Fed's; BoI phrasings that have no Fed analogue", False),
    ("  are covered only by the 'judgment' terms.", False),
    ("- A ratio of cue counts, not a semantic model: it cannot read negation or context the way", False),
    ("  the sentence-tone pass (other dashboard tab) or Doh-Song-Yang's embedding model can.", False),
    ("- Counts are corpus data, not spreadsheet formulas (LibreOffice recalculation is", False),
    ("  unavailable on this machine); the arithmetic is spelled out in the 'Calculation' column.", False),
]
for text, is_head in meth:
    ws3.append([text])
    ws3.cell(row=ws3.max_row, column=1).font = (
        Font(name="Arial", size=11, bold=True, color="1F4E79") if is_head else Font(name="Arial", size=10))
ws3.column_dimensions["A"].width = 108

wb.save(OUT)
print("wrote", OUT, "|", len(LEXICON), "terms |", len(scores), "announcements")
