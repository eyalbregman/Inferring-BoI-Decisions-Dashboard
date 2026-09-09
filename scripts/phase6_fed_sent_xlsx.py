# -*- coding: utf-8 -*-
"""results/fed_sentence_tone.xlsx -- Fed-anchored hawkish/dovish tone, sentence by sentence.

Sheet 1  Announcement scores : one row per raw BoI announcement, score = (H - D) / total
Sheet 2  All sentences       : every sentence with its H / D / N label
Sheet 3  Method
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'results', 'fed_sentence_tone.xlsx')

scores = json.load(open(os.path.join(ROOT, 'results', 'fed_sentence_scores.json'), encoding='utf-8'))
full = json.load(open(os.path.join(ROOT, 'results', 'fed_sentence_labels_full.json'), encoding='utf-8'))

BASE = Font(name="Arial", size=10)
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F4E79")
FILL_H = PatternFill("solid", fgColor="FBE4E1")
FILL_D = PatternFill("solid", fgColor="DEEAF2")
FILL_N = PatternFill("solid", fgColor="F2F2F2")
POS = PatternFill("solid", fgColor="F4C7C3")
NEG = PatternFill("solid", fgColor="BDD7EE")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def head(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADF; cell.fill = HEADFILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    ws.row_dimensions[1].height = 30

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Sheet 1
ws = wb.active
ws.title = "Announcement scores"
ws.append(["Date", "Decision (reference only)", "Hawkish sentences (H)", "Dovish sentences (D)",
           "Neutral sentences (N)", "Total sentences", "Score  (H - D) / Total", "Calculation",
           "Lexicon-hit score (other method)"])
for i, r in enumerate(scores, start=2):
    ws.append([r['date'], r['decision'], r['H'], r['D'], r['N'], r['n'], round(r['score'], 4),
               f"({r['H']} - {r['D']}) / {r['n']} = {r['score']:+.3f}",
               round(r['lex_score'], 3) if r.get('lex_score') is not None else None])
    for c in range(1, 10):
        cell = ws.cell(row=i, column=c); cell.font = BASE; cell.border = BORDER
        if c in (3, 4, 5, 6, 9):
            cell.alignment = Alignment(horizontal="center")
    sc = ws.cell(row=i, column=7)
    sc.number_format = "+0.000;-0.000;0.000"; sc.font = Font(name="Arial", size=10, bold=True)
    sc.fill = POS if r['score'] > 0 else (NEG if r['score'] < 0 else FILL_N)
    ws.cell(row=i, column=3).fill = FILL_H
    ws.cell(row=i, column=4).fill = FILL_D
    ws.cell(row=i, column=5).fill = FILL_N
    lc = ws.cell(row=i, column=9)
    if r.get('lex_score') is not None:
        lc.number_format = "+0.000;-0.000;0.000"
for i, w in enumerate([12, 22, 15, 14, 14, 12, 16, 24, 20], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
head(ws, 9)

sH = sum(r['H'] for r in scores); sD = sum(r['D'] for r in scores)
sN = sum(r['N'] for r in scores); sT = sH + sD + sN
b = len(scores) + 3
ws.cell(row=b, column=1, value="Overall (all announcements)").font = Font(name="Arial", size=10, bold=True)
for col, val in ((3, sH), (4, sD), (5, sN), (6, sT)):
    ws.cell(row=b, column=col, value=val).font = Font(name="Arial", size=10, bold=True)
sc = ws.cell(row=b, column=7, value=round((sH - sD) / sT, 4))
sc.font = Font(name="Arial", size=10, bold=True); sc.number_format = "+0.000;-0.000;0.000"
ws.cell(row=b, column=8, value=f"({sH} - {sD}) / {sT} = {(sH-sD)/sT:+.3f}").font = BASE
for j, (lab, dec) in enumerate([("mean score, hikes", "raise"), ("mean score, holds", "maintain"),
                                ("mean score, cuts", "lower")], start=1):
    v = [r['score'] for r in scores if r['decision'] == dec]
    ws.cell(row=b + j, column=1, value=lab).font = BASE
    mc = ws.cell(row=b + j, column=7, value=round(sum(v) / len(v), 4))
    mc.font = BASE; mc.number_format = "+0.000;-0.000;0.000"

# ---------------------------------------------------------------- Sheet 2
ws2 = wb.create_sheet("All sentences")
ws2.append(["Date", "#", "Tone", "Sentence"])
for r in full:
    ws2.append([r['id'], r['i'], r['label'], r['sentence']])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BASE
    row[1].alignment = Alignment(horizontal="center")
    row[2].alignment = Alignment(horizontal="center")
    row[2].font = Font(name="Arial", size=10, bold=True)
    row[2].fill = {"H": FILL_H, "D": FILL_D, "N": FILL_N}.get(row[2].value, FILL_N)
    row[3].alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([12, 5, 7, 120], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
head(ws2, 4)

# ---------------------------------------------------------------- Sheet 3
ws3 = wb.create_sheet("Method")
NH = sum(r['H'] for r in scores); ND = sum(r['D'] for r in scores)
NN = sum(r['N'] for r in scores); NT = NH + ND + NN
mh = sum(r['score'] for r in scores if r['decision'] == 'raise') / max(1, sum(1 for r in scores if r['decision'] == 'raise'))
mm = sum(r['score'] for r in scores if r['decision'] == 'maintain') / max(1, sum(1 for r in scores if r['decision'] == 'maintain'))
ml = sum(r['score'] for r in scores if r['decision'] == 'lower') / max(1, sum(1 for r in scores if r['decision'] == 'lower'))
meth = [
    ("Fed-anchored hawkish / dovish tone, sentence by sentence", True),
    ("", False),
    ("The basis", True),
    ("The US Federal Reserve staff drafts, for every FOMC meeting, a fully DOVISH version of", False),
    ("its statement (Alternative A) and a fully HAWKISH version (Alternative C / D). We pooled", False),
    ("75 meetings of those drafts and learned which content and phrasings are hawkish vs. dovish", False),
    ("in the Fed's own extremes (see the 'hawkish_dovish_fed_anchored.xlsx' lexicon). Here we", False),
    ("apply that basis at the SENTENCE level: each sentence of every BoI announcement is read", False),
    ("and judged - would this sentence belong in the Fed's hawkish draft, its dovish draft, or", False),
    ("neither? This is about direction / content, not about how confident the phrasing sounds", False),
    ("(that is the separate 'Sentence tone' measure).", False),
    ("", False),
    ("Corpus", True),
    ("62 Bank of Israel English announcements, Nov 2018 - today, RAW text: the headline and the", False),
    ("decision sentence are KEPT (this measure is decision-inclusive, like the paper). Navigation,", False),
    ("dates, figure captions, Hebrew and the procedural closing are removed; the repeated headline", False),
    ("is collapsed to one. %d sentences in total." % NT, False),
    ("", False),
    ("Labels", True),
    ("H  hawkish - fits the Fed's Alt C/D: the decision/guidance is to raise or tighten; inflation", False),
    ("   is high / above target / broad / accelerating / a persistent risk; the economy is strong,", False),
    ("   the labour market tight / at full employment; the shekel weakened; upside inflation risk.", False),
    ("D  dovish - fits the Fed's Alt A: the decision/guidance is to lower / 'accommodative' /", False),
    ("   'support activity' / patience / low-for-long; inflation is low / below target / moderating;", False),
    ("   the economy is soft / weak / slowing / contracting; downside risks; global drag; the", False),
    ("   shekel strengthened.", False),
    ("N  neutral - plain data or description with no clear lean, procedural, balanced, or a", False),
    ("   neutrally-stated forecast. The decision sentence for a HOLD is N.", False),
    ("", False),
    ("Every sentence and its label is on the 'All sentences' sheet.", False),
    ("", False),
    ("Score", True),
    ("For each announcement:   score = (H - D) / total sentences,   -1 to +1.", False),
    ("The 'Calculation' column shows the arithmetic. The last column repeats the lexicon-hit", False),
    ("score (counting matched hawk/dove terms) from hawkish_dovish_fed_anchored.xlsx for comparison.", False),
    ("", False),
    ("What it shows", True),
    ("Corpus mix H %.1f%% / D %.1f%% / N %.1f%%.  Mean score by decision:  hikes %+.2f,  holds %+.2f,  cuts %+.2f."
     % (100*NH/NT, 100*ND/NT, 100*NN/NT, mh, mm, ml), False),
    ("Classification: one pass by Claude (Sonnet) over every sentence, batched by announcement.", False),
]
for text, is_head in meth:
    ws3.append([text])
    ws3.cell(row=ws3.max_row, column=1).font = (
        Font(name="Arial", size=11, bold=True, color="1F4E79") if is_head else Font(name="Arial", size=10))
ws3.column_dimensions["A"].width = 110

wb.save(OUT)
print("wrote", OUT, "|", len(scores), "announcements |", len(full), "sentences")
