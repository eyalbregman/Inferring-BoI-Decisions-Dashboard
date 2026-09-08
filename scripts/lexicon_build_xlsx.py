# -*- coding: utf-8 -*-
"""Build results/hawkish_dovish_lexicon.xlsx from the hand-curated lexicon + corpus stats.

Sheets:
  Lexicon             - curated terms sorted Dimension > Direction > Term, colour-coded
  All candidate terms - every n-gram found >=4x in >=3 announcements (for auditing / additions)
  Method              - how this was built
"""
import os, re, glob, csv, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from lexicon_data import LEXICON

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'results', 'hawkish_dovish_lexicon.xlsx')

DIR_NAME = {"H": "Hawkish (strong / confident)", "D": "Dovish (soft / hedged)", "N": "Neutral (not scored)"}
DIR_ORDER = {"H": 0, "D": 1, "N": 2}

# ---------------------------------------------------------------- corpus
DROP_LINE = re.compile(
    r'^(Home Page|Communication and [Pp]ublications|Press Releases|Back|Share:?|'
    r'The Monetary Committee decides on|to (lower|raise|increase|reduce|leave|keep) the interest|'
    r'\d{1,2}/\d{1,2}/\d{2,4}|To view this|For the file|This page was last updated|'
    r'The minutes of the monetary discussions|The next decision regarding the interest rate)', re.I)

def body(text):
    out = []
    for ln in text.splitlines():
        ln = ln.strip()
        if not ln or DROP_LINE.match(ln):
            continue
        if re.fullmatch(r'[\u0590-\u05FF\s%\d.,-]+', ln):
            continue
        out.append(ln)
    t = " ".join(out)
    t = re.sub(r'\s+', ' ', t)
    return t

def split_sents(t):
    return [s.strip() for s in re.split(r'(?<=[.;:])\s+(?=[A-Z“"\u2018])', t) if len(s.strip()) > 15]

docs = sorted(glob.glob(os.path.join(ROOT, 'data', 'raw', '*.txt')))
bodies, sent_lists = [], []
for d in docs:
    b = body(open(d, encoding='utf-8', errors='replace').read())
    bodies.append(b.lower())
    sent_lists.append(split_sents(b))
N = len(docs)

_GAP = r"(?:'s|’s)?[\s\-]+"   # word gap: space or hyphen, tolerate possessive 's

def term_regex(term):
    """Forgiving match: ' ... ' = a within-sentence gap; words may be separated by
    spaces or hyphens; 'percent' may be preceded by a number ('by percent' -> 'by 5.2 percent')."""
    def one(seg):
        words = re.split(r'[\s\-]+', seg.strip())
        p = _GAP.join(re.escape(w) for w in words)
        p = p.replace('percent', r'(?:[\d.,]+\s*)?percent')
        return r'\b' + p + r'\b'          # whole-word: 'low' != 'below', 'high' != 'higher'
    return re.compile(r'[^.;:]*?'.join(one(s) for s in term.split(' ... ')), re.I)

def stats(term):
    rx = term_regex(term)
    total = docf = 0
    example = ""
    for i in range(N):
        c = len(rx.findall(bodies[i]))
        if c:
            total += c
            docf += 1
        if not example:
            cands = [s for s in sent_lists[i] if rx.search(s)]
            cands = [s for s in cands if 35 <= len(s) <= 240] or cands
            if cands:
                example = min(cands, key=len)
    example = re.sub(r'\s+', ' ', example).strip().replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
    return total, docf, example

# ---------------------------------------------------------------- assemble curated rows
rows = []
for term, dim, d, logic in LEXICON:
    total, docf, ex = stats(term)
    typ = "phrase" if (" " in term or "-" in term) else "word"
    rows.append([term, typ, dim, DIR_NAME[d], logic, ex, total, docf, d])

rows.sort(key=lambda r: (r[2].lower(), DIR_ORDER[r[8]], r[0].lower()))

missing = [r for r in rows if r[6] == 0]

# ---------------------------------------------------------------- workbook
wb = openpyxl.Workbook()
BASE = Font(name="Arial", size=10)
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F4E79")
FILL_H = PatternFill("solid", fgColor="FBE4E1")   # light red  - strong / confident
FILL_D = PatternFill("solid", fgColor="DEEAF2")   # light blue - soft / hedged
FILL_C = PatternFill("solid", fgColor="F0F0F0")   # light grey - neutral
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADF; cell.fill = HEADFILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    ws.row_dimensions[1].height = 28

# --- Sheet 1: Lexicon
ws = wb.active
ws.title = "Lexicon"
HEAD = ["Term", "Type", "Rhetorical category", "Direction",
        "Why it reads as strong / hedged", "Example sentence from the 62 announcements",
        "Occurrences", "In N of 62"]
ws.append(HEAD)
for r in rows:
    ws.append(r[:8])
for i, r in enumerate(rows, start=2):
    fill = {"H": FILL_H, "D": FILL_D, "N": FILL_C}[r[8]]
    for c in range(1, 9):
        cell = ws.cell(row=i, column=c)
        cell.font = BASE; cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (1, 3, 5, 6)))
        if c == 4:
            cell.fill = fill
            cell.font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=i, column=7).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=i, column=8).alignment = Alignment(horizontal="center", vertical="top")
widths = [26, 9, 24, 27, 52, 66, 12, 11]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
style_header(ws, 8)

# --- Sheet 2: All candidate terms (frequency dump for auditing)
ws2 = wb.create_sheet("All candidate terms")
ws2.append(["Term", "Words", "Occurrences", "In N of 62", "Example sentence",
            "Direction H/D/N (fill in)", "Rhetorical category (fill in)"])
freq_csv = os.path.join(ROOT, 'data', 'ngram_freq.csv')
labelled = {r[0].lower() for r in rows}
with open(freq_csv, encoding='utf-8') as f:
    for term, nlen, tot, df, ex in list(csv.reader(f))[1:]:
        mark = "  (already in Lexicon)" if term.lower() in labelled else ""
        ws2.append([term, int(nlen), int(tot), int(df),
                    re.sub(r'\s+', ' ', ex).strip()[:240] + mark, "", ""])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BASE
    row[4].alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([34, 8, 12, 11, 70, 16, 18], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
style_header(ws2, 7)

# --- Sheet 3: Method
ws3 = wb.create_sheet("Method")
method = [
    ("Hawkish / Dovish lexicon - Bank of Israel interest-rate announcements", True),
    ("", False),
    ("Source corpus", True),
    ("62 Bank of Israel English interest-rate announcements, 26 Nov 2018 - 6 Jul 2026", False),
    ("(data/raw/*.txt in this repo). Navigation, the headline, the decision sentence,", False),
    ("the date stamp and Hebrew snippets are removed before counting; everything else is kept.", False),
    ("", False),
    ("What this lexicon measures", True),
    ("The STRENGTH AND CONFIDENCE OF THE LANGUAGE - not the economic content, not the policy decision.", False),
    ("", False),
    ("Hawkish (H) = strong / confident / emphatic phrasing. The sentence commits to a claim,", False),
    ("              amplifies it, or asserts it flatly:", False),
    ("              'rose sharply', 'significant', 'remains tight', 'will', 'record high', 'in particular'.", False),
    ("Dovish  (D) = soft / hedged / tentative phrasing. The sentence qualifies, downplays or doubts", False),
    ("              the claim:", False),
    ("              'rose slightly', 'moderate', 'may', 'appears', 'around 3 percent', 'uncertainty', 'so far'.", False),
    ("Neutral (N) = a plain factual verb or noun with no rhetorical charge on its own", False),
    ("              ('increased', 'declined', 'stable', 'remains'). Listed only to mark it as deliberately unscored.", False),
    ("", False),
    ("The label is subject-independent.", True),
    ("'increased sharply' is strong language whether the subject is inflation, unemployment, the shekel", False),
    ("or bond yields - and whether that month's decision was lower, maintain or raise. A bare 'increased'", False),
    ("is neutral; the signal lives in the modifier ('sharply' vs 'slightly'), the modal ('will' vs 'may')", False),
    ("and the framing ('remains' vs 'appears'), never in the noun.", False),
    ("", False),
    ("Rhetorical categories used", True),
    ("Magnitude - amplifier / dampener   sharp, significant, rapid  vs  slight, moderate, gradual", False),
    ("Firmness / persistence             remains tight, persistent, entrenched, sticky", False),
    ("Assertion / commitment             will, must, determined, emphasized, as necessary", False),
    ("Emphasis marker                    in particular, particularly, in fact", False),
    ("Level / state                      high, elevated, tight, full employment (flat, unqualified)", False),
    ("Hedge / qualifier                  expected to, projected, broadly, largely, to some extent", False),
    ("Tentativeness - modal              may, could, possible, appears, seems, likely", False),
    ("Approximation                      around, approximately, close to, in the range of", False),
    ("Uncertainty                        uncertainty, volatile, mixed, in opposite directions, so far", False),
    ("Caution / patience                 cautious, monitor, closely (wait-and-see - defers judgement)", False),
    ("", False),
    ("Columns", True),
    ("Occurrences  = total times the term appears across the 62 announcements (case-insensitive; ' ... ' = a gap within one sentence).", False),
    ("In N of 62   = number of separate announcements that contain the term at least once.", False),
    ("These counts are read from the corpus by scripts/lexicon_build_xlsx.py; they are data, not spreadsheet formulas.", False),
    ("", False),
    ("'All candidate terms' sheet", True),
    ("Every 1-3 word phrase that appears at least 4 times in at least 3 announcements (~3,400 rows),", False),
    ("unlabelled, so you can scan for anything the curated Lexicon missed and fill in Direction / Category.", False),
    ("", False),
    ("Next step", True),
    ("Turn the reviewed Lexicon into a machine-readable dictionary (JSON / regex) and use it to score", False),
    ("each blinded rationale - a transparent 'confidence index' (strong-minus-hedged, per 1,000 words)", False),
    ("to compare against the LLM classifier and against the actual decisions.", False),
]
for text, is_head in method:
    ws3.append([text])
    cell = ws3.cell(row=ws3.max_row, column=1)
    cell.font = Font(name="Arial", size=11, bold=True, color="1F4E79") if is_head else Font(name="Arial", size=10)
ws3.column_dimensions["A"].width = 120

wb.save(OUT)
print("wrote", OUT)
print("curated terms:", len(rows),
      "| Hawkish:", sum(1 for r in rows if r[8] == 'H'),
      "Dovish:", sum(1 for r in rows if r[8] == 'D'),
      "Neutral:", sum(1 for r in rows if r[8] == 'N'))
if missing:
    print("\nNOT FOUND in corpus (fix the term string):")
    for r in missing:
        print("  -", r[0])
