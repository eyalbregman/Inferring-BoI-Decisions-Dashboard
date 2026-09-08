# -*- coding: utf-8 -*-
"""results/sentence_tone.xlsx  -- sentence-level hawkish/dovish tone of each announcement.

Sheet 1  Announcement tone scores : one row per announcement, score = (H - D) / total
Sheet 2  All sentences            : every sentence with its H/D/N label
Sheet 3  Method
"""
import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'results', 'sentence_tone.xlsx')

scores = json.load(open(os.path.join(ROOT, 'results', 'sentence_scores.json'), encoding='utf-8'))
full = json.load(open(os.path.join(ROOT, 'results', 'sentence_labels_full.json'), encoding='utf-8'))

BASE = Font(name="Arial", size=10)
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="1F4E79")
FILL_H = PatternFill("solid", fgColor="FBE4E1")
FILL_D = PatternFill("solid", fgColor="DEEAF2")
FILL_N = PatternFill("solid", fgColor="F2F2F2")
POS = PatternFill("solid", fgColor="F4C7C3")   # hawkish score
NEG = PatternFill("solid", fgColor="BDD7EE")   # dovish score
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def head(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADF; cell.fill = HEADFILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"
    ws.row_dimensions[1].height = 30

wb = openpyxl.Workbook()

# ---------------------------------------------------------------- Sheet 1
ws = wb.active
ws.title = "Announcement tone scores"
H1 = ["Date", "Decision (reference only - not used)", "Hawkish sentences (H)",
      "Dovish sentences (D)", "Neutral sentences (N)", "Total sentences",
      "Tone score  (H - D) / Total", "Calculation"]
ws.append(H1)
for i, r in enumerate(scores, start=2):
    ws.append([r['date'], r['decision'], r['H'], r['D'], r['N'], r['n'],
               round(r['score'], 4),
               f"({r['H']} - {r['D']}) / {r['n']} = {r['score']:+.3f}"])
    for c in range(1, 9):
        cell = ws.cell(row=i, column=c); cell.font = BASE; cell.border = BORDER
        if c in (3, 4, 5, 6):
            cell.alignment = Alignment(horizontal="center")
    sc = ws.cell(row=i, column=7)
    sc.number_format = "+0.000;-0.000;0.000"
    sc.font = Font(name="Arial", size=10, bold=True)
    sc.fill = POS if r['score'] > 0 else (NEG if r['score'] < 0 else FILL_N)
    ws.cell(row=i, column=3).fill = FILL_H
    ws.cell(row=i, column=4).fill = FILL_D
    ws.cell(row=i, column=5).fill = FILL_N
for i, w in enumerate([12, 26, 15, 14, 14, 12, 16, 24], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
head(ws, 8)

# summary block below the table (static totals - LibreOffice recalc is unavailable on this host)
sH = sum(r['H'] for r in scores); sD = sum(r['D'] for r in scores)
sN = sum(r['N'] for r in scores); sT = sH + sD + sN
base = len(scores) + 3
ws.cell(row=base, column=1, value="Overall (all announcements)").font = Font(name="Arial", size=10, bold=True)
for col, val in ((3, sH), (4, sD), (5, sN), (6, sT)):
    ws.cell(row=base, column=col, value=val).font = Font(name="Arial", size=10, bold=True)
sc = ws.cell(row=base, column=7, value=round((sH - sD) / sT, 4))
sc.font = Font(name="Arial", size=10, bold=True); sc.number_format = "+0.000;-0.000;0.000"
ws.cell(row=base, column=8, value=f"({sH} - {sD}) / {sT} = {(sH-sD)/sT:+.3f}").font = BASE
mean = sum(r['score'] for r in scores) / len(scores)
ws.cell(row=base + 1, column=1, value="Mean of the 62 announcement scores").font = BASE
m = ws.cell(row=base + 1, column=7, value=round(mean, 4))
m.font = BASE; m.number_format = "+0.000;-0.000;0.000"

# ---------------------------------------------------------------- Sheet 2
ws2 = wb.create_sheet("All sentences")
ws2.append(["Date", "#", "Tone", "Sentence"])
for r in full:
    ws2.append([r['id'], r['i'], r['label'], r['sentence']])
for row in ws2.iter_rows(min_row=2):
    for cell in row:
        cell.font = BASE
    row[1].alignment = Alignment(horizontal="center")
    row[2].alignment = Alignment(horizontal="center")
    row[2].font = Font(name="Arial", size=10, bold=True)
    row[2].fill = {"H": FILL_H, "D": FILL_D, "N": FILL_N}.get(row[2].value, FILL_N)
    row[3].alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([12, 5, 7, 120], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
head(ws2, 4)

# ---------------------------------------------------------------- Sheet 3
ws3 = wb.create_sheet("Method")
NH = sum(r['H'] for r in scores); ND = sum(r['D'] for r in scores)
NN = sum(r['N'] for r in scores); NT = NH + ND + NN
meth = [
    ("Sentence-level hawkish / dovish tone", True),
    ("", False),
    ("Why sentences, not words", True),
    ("A word carries different weight depending on its sentence ('sharp' in 'a sharp rise' vs.", False),
    ("'less sharp than expected'). So each announcement is split into sentences and every", False),
    ("sentence is read and classified as a whole.", False),
    ("", False),
    ("Corpus", True),
    ("62 Bank of Israel English announcements, Nov 2018 - today, blinded text (the decision", False),
    ("sentence, rate figure and dates already removed). Navigation, figure captions, Hebrew and", False),
    ("the procedural closing paragraph are stripped. %d sentences in total." % NT, False),
    ("", False),
    ("Labels - by the STRENGTH AND CONFIDENCE of the wording, not the decision", True),
    ("H  strong / confident / emphatic : intensifiers (sharp, significant, marked, rapid, steep,", False),
    ("   strong, robust, broad-based, record, accelerated), flat declaratives ('remains tight',", False),
    ("   'is very low'), firm persistence, categorical 'will / must', emphasis ('in particular').", False),
    ("D  soft / hedged / tentative : dampeners (slight, moderate, gradual, somewhat, relatively,", False),
    ("   limited), epistemic hedges (may, could, appears, seems, likely, expected to, projected,", False),
    ("   assessed, estimated), approximation (around, close to), uncertainty (volatile, mixed,", False),
    ("   'so far'), wait-and-see (cautious, will monitor).", False),
    ("N  neutral : plain data, background, definitions, or strong and hedged elements balanced.", False),
    ("", False),
    ("The label ignores whether the news is economically good or bad and ignores any implied", False),
    ("rate direction. 'Inflation fell sharply' = H (because 'sharply'); 'inflation may have eased", False),
    ("somewhat' = D; 'inflation was 3.0 percent' = N.", False),
    ("", False),
    ("Every sentence and its label is on the 'All sentences' sheet.", False),
    ("", False),
    ("Score", True),
    ("For each announcement:   score = (H - D) / total sentences", False),
    ("  -1.0 = every sentence dovish/hedged      +1.0 = every sentence hawkish/confident", False),
    ("   0.0 = equal H and D, or all neutral", False),
    ("Neutral sentences are in the denominator, so a mostly factual announcement scores near 0.", False),
    ("The 'Calculation' column shows the arithmetic (H, D and total) behind every score.", False),
    ("", False),
    ("Corpus-wide: H = %.1f%%, D = %.1f%%, N = %.1f%% of all %d sentences." % (
        100*NH/NT, 100*ND/NT, 100*NN/NT, NT), False),
    ("Classification: one pass by Claude (Sonnet) over every sentence, batched by announcement.", False),
]
for text, is_head in meth:
    ws3.append([text])
    ws3.cell(row=ws3.max_row, column=1).font = (
        Font(name="Arial", size=11, bold=True, color="1F4E79") if is_head else Font(name="Arial", size=10))
ws3.column_dimensions["A"].width = 110

wb.save(OUT)
print("wrote", OUT)
print("rows:", len(scores), "| sentences:", len(full))
